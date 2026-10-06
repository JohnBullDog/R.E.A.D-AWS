import io
import sys
import types
import zipfile

import docx
import pytest

from read.chunk import build_chunks
from read.extract import Unsupported, extract, sniff


def make_docx() -> bytes:
    d = docx.Document()
    d.add_heading("Recommendation 1", level=1)
    d.add_paragraph("Teach students to recognize the sounds in words.")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Grade", "Skill"
    t.cell(1, 0).text, t.cell(1, 1).text = "K", "Blending"
    d.add_paragraph("   ")
    d.add_heading("Recommendation 2", level=1)
    d.add_paragraph("Read connected text daily.")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def test_sniff_by_content_not_extension():
    assert sniff(b"%PDF-1.7 ...") == "pdf"
    assert sniff(make_docx()) == "docx"
    assert sniff(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1rest") == "doc"
    assert sniff(b"plain words") == "txt"


def test_non_word_zip_rejected():
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("hello.txt", "hi")
    with pytest.raises(Unsupported):
        sniff(buf.getvalue())


def test_doc_rejected_while_libreoffice_deferred():
    with pytest.raises(Unsupported, match="docx"):
        extract(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 100)


def test_txt_normalizes_newlines_and_bom():
    data = "﻿Line one\r\nLine two\rLine three — café".encode()
    ex = extract(data)
    assert ex.text == "Line one\nLine two\nLine three — café"
    assert "\r" not in ex.text


def test_txt_detects_windows_1252():
    para = (
        "Teachers “model” each sound — then students blend them. "
        "Students can’t skip this step; it’s the foundation for decoding. "
    )
    data = (para * 4).encode("cp1252")
    text = extract(data).text
    assert "“model”" in text and "—" in text and "can’t" in text


def test_binary_rejected():
    with pytest.raises(Unsupported):
        extract(b"\x00\x01\x02\x03" * 200)


def test_docx_headings_tables_order():
    ex = extract(make_docx())
    assert ex.text.split("\n\n") == [
        "Recommendation 1",
        "Teach students to recognize the sounds in words.",
        "Grade | Skill",
        "K | Blending",
        "Recommendation 2",
        "Read connected text daily.",
    ]
    assert ex.headings == {"Recommendation 1", "Recommendation 2"}
    sections, _ = build_chunks("w", "v", ex.text, ex.headings)
    assert [s["section_path"] for s in sections] == ["Recommendation 1", "Recommendation 2"]


class FakePage:
    def __init__(self, text):
        self.text = text

    def extract_text(self):
        return self.text


class FakePdf:
    def __init__(self, pages):
        self.pages = [FakePage(t) for t in pages]

    def __enter__(self):
        return self

    def __exit__(self, *a):
        return False


def fake_pdfplumber(monkeypatch, pages):
    mod = types.SimpleNamespace(open=lambda _f: FakePdf(pages))
    monkeypatch.setitem(sys.modules, "pdfplumber", mod)


def test_pdf_page_offsets(monkeypatch):
    p1, p2 = "A" * 300 + "\r\nmore", "B" * 300
    fake_pdfplumber(monkeypatch, [p1, p2])
    ex = extract(b"%PDF-1.7")
    assert ex.page_starts[0] == 0
    assert ex.text[ex.page_starts[1] :].startswith("B")
    assert "\r" not in ex.text
    _, passages = build_chunks("w", "v", ex.text, page_starts=ex.page_starts, target=1)
    assert [p["page"] for p in passages] == [1, 2]


def test_scanned_pdf_rejected(monkeypatch):
    fake_pdfplumber(monkeypatch, ["", "  ", None])
    with pytest.raises(Unsupported, match="Textract"):
        extract(b"%PDF-1.7")
