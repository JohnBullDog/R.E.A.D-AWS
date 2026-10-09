import io
import os
import re
import time
from pathlib import Path

import pytest

from read import review as R
from read import review_run
from read.extract import extract
from read.store import TempStore

SHORT = (
    "Objective: students blend three sounds to read short words.\n\n"
    "Teacher says c-a-t slowly and students say cat. Repeat with map, sit, and hop.\n\n"
    "Exit ticket: each student reads five words aloud."
)
PARTS = R.material_parts(SHORT)
SECTION = {
    "id": "sec1",
    "title": "Whole material",
    "parts": ["M1", "M2", "M3"],
    "owned": ["M1", "M2", "M3"],
}
RESEARCH = {
    "S1": (
        "Teach students to blend individual sounds into words using short daily routines "
        "with feedback."
    ),
    "S2": "Check student progress often with short assessments of word reading.",
}
TEXTS = {p["id"]: p["text"] for p in PARTS} | RESEARCH


def sent(text, material=(), cites=()):
    return {"text": text, "material": list(material), "cites": list(cites)}


def section_out(**kw):
    out = {
        "findings": [
            {
                "question_id": "phonological-awareness",
                "seen": "yes",
                "observation": [sent("The lesson models blending with three short words.", ["M2"])],
                "suggestions": [
                    sent("Add brief corrective feedback during blending.", cites=["S1"])
                ],
            }
        ],
    }
    out.update(kw)
    return out


def validate(out):
    return R.validate_section(out, SECTION, {"phonological-awareness"}, TEXTS, RESEARCH)


# ---------------- checklist ----------------


def test_shipped_checklist_is_valid_two_lists_and_placeholder():
    data = R.load_checklist(Path("rubric/checklist.json"))
    assert R.check_checklist(data) == []
    assert data["status"] == "DEVELOPMENT PLACEHOLDER" and "rewritten by Addison" in data["note"]
    a, b = R.split_checklist(data["criteria"], "K-12")
    assert len(a) == 7 and len(b) == 7
    assert {q["criterion_id"] for q in b if q["core"]} == {
        "doc-objective",
        "doc-alignment",
        "doc-check-learning",
    }


def test_split_checklist_by_grade():
    data = R.load_checklist(Path("rubric/checklist.json"))
    a, _ = R.split_checklist(data["criteria"], "4")
    assert "phonological-awareness" not in {q["criterion_id"] for q in a}  # K-1 only


def test_check_checklist_catches_scope_and_combine():
    bad = {
        "criteria": [
            {
                "criterion_id": "x",
                "component": "c",
                "grade_band": "K",
                "question": "q",
                "search_query": "s",
                "scope": "document",
                "combine": "sometimes",
            },
            {
                "criterion_id": "y",
                "component": "c",
                "grade_band": "K",
                "question": "q",
                "search_query": "s",
                "scope": "page",
            },
        ]
    }
    probs = R.check_checklist(bad)
    assert any("combine must be one of" in p for p in probs)
    assert any("scope must be section or document" in p for p in probs)


# ---------------- sections ----------------


def para(n, word="word"):
    return " ".join([word] * (n - 1)) + " end."


def test_short_material_is_one_section():
    assert R.make_sections(PARTS) == [
        {
            "id": "sec1",
            "title": "Whole material",
            "parts": ["M1", "M2", "M3"],
            "owned": ["M1", "M2", "M3"],
        }
    ]


def test_headed_material_splits_at_headings_caps_and_merges_tiny_groups():
    text = "\n\n".join(
        [
            "Unit plan",
            "Day 1: Blending",
            para(400),
            para(400),
            para(300),
            "Day 2: Segmenting",
            para(500),
            "Wrap-up",
            para(200),
        ]
    )
    parts = R.material_parts(text)
    secs = R.make_sections(parts)
    titles = [s["title"] for s in secs]
    assert titles[0] == "Day 1: Blending"  # a lone heading merges into the next group
    assert "Day 1: Blending (cont.)" in titles  # 1,100 words capped at 800
    assert titles[-2:] == ["Day 2: Segmenting", "Wrap-up"]
    assert all(s["owned"] == s["parts"] for s in secs)
    flat = [m for s in secs for m in s["parts"]]
    assert flat == [p["id"] for p in parts]  # every part once, in order


def test_unstructured_material_uses_sliding_windows_with_overlap():
    text = "\n\n".join(para(80) for _ in range(30))  # 2,400 words, no headings
    parts = R.material_parts(text)
    secs = R.make_sections(parts)
    assert len(secs) > 1 and secs[0]["title"] == "Part 1"
    second = secs[1]
    overlap = [m for m in second["parts"] if m not in second["owned"]]
    assert overlap and set(overlap) <= set(secs[0]["parts"])
    owned = [m for s in secs for m in s["owned"]]
    assert owned == [p["id"] for p in parts]  # each part owned exactly once


def test_dedupe_overlap_drops_findings_only_about_overlap():
    sec = {"parts": ["M1", "M2", "M3"], "owned": ["M2", "M3"]}
    out = {
        "findings": [
            {"observation": [sent("a", ["M1"])]},
            {"observation": [sent("b", ["M1", "M2"])]},
        ],
        "doc": [],
    }
    assert len(R.dedupe_overlap(out, sec)["findings"]) == 1


# ---------------- triage and goal ----------------


def test_triage_thresholds():
    unit = [1.0, 0.0]
    t = R.triage(unit, {"a": [0.9, 0.436], "b": [0.3, 0.954], "c": [0.1, 0.995]})
    assert (t["a"]["label"], t["b"]["label"], t["c"]["label"]) == ("in", "maybe", "out")


def test_clean_goal_forces_core_and_keeps_reasons():
    qs = [
        {"criterion_id": "doc-objective", "question": "q1", "core": True},
        {"criterion_id": "doc-support", "question": "q2", "core": False},
        {"criterion_id": "doc-sequence", "question": "q3", "core": False},
    ]
    g = R.clean_goal(
        {
            "grade_band": "Grade 1",
            "objective": "o",
            "doc_questions": [
                {"id": "doc-objective", "relevant": False, "reason": "x"},
                {"id": "doc-support", "relevant": True, "reason": "has groups"},
            ],
        },
        qs,
    )
    assert g["grade_band"] == "1"
    rel = {d["id"]: d["relevant"] for d in g["doc_questions"]}
    assert rel == {"doc-objective": True, "doc-support": True, "doc-sequence": False}
    assert g["doc_questions"][1]["reason"] == "has groups"


# ---------------- section validation ----------------


def test_valid_section_passes():
    assert validate(section_out()) == []


def test_section_rejects_unknown_ids_and_outside_material():
    f = section_out()["findings"][0]
    out = section_out(
        findings=[
            {"question_id": "nope", "seen": "yes", "observation": [], "suggestions": []},
            {**f, "observation": [sent("It blends map and sit.", ["M9"])]},
        ]
    )
    probs = validate(out)
    assert any("must use one of the section question ids" in p for p in probs)
    assert any("outside this section" in p for p in probs)


def test_section_finding_rules():
    f = section_out()["findings"][0]
    probs = validate(section_out(findings=[{**f, "observation": [sent("It blends.")]}]))
    assert any("must point to the material parts" in p for p in probs)
    probs = validate(section_out(findings=[{**f, "suggestions": [sent("Do more.")]}]))
    assert any("must cite a research section" in p for p in probs)


def test_reused_material_wording_is_a_quote():
    s = sent(
        "It says students blend three sounds to read short words, which fits the goal.", ["M1"]
    )
    out = section_out(findings=[{**section_out()["findings"][0], "observation": [s]}])
    assert validate(out) == []
    parts = R.show([s], TEXTS)[0]["parts"]
    assert any(p.get("quote") == "students blend three sounds to read short words" for p in parts)


# ---------------- combine and summary ----------------


def test_combine_status_rules():
    assert R.combine_status("any", ["none", "partial", "none"]) == "partly"
    assert R.combine_status("any", ["none", "full"]) == "met"
    assert R.combine_status("any", ["none", "none"]) == "missing"
    assert R.combine_status("all", ["full", "full"]) == "met"
    assert R.combine_status("all", ["full", "none"]) == "partly"
    assert R.combine_status("all", ["none"]) == "missing"
    assert R.combine_status("judge", ["full"]) is None


def test_validate_combine_enforces_fixed_and_missing_rules():
    per = {
        "a": {"status": "met", "answers": [{"answer": "full"}]},
        "b": {"status": None, "answers": [{"answer": "partial"}]},
    }
    out = {
        "verdicts": [
            {
                "question_id": "a",
                "explicit": True,
                "verdict": "partly",
                "observation": [sent("x", ["M1"])],
                "suggestions": [sent("y", cites=["S1"])],
            },
            {
                "question_id": "b",
                "verdict": "missing",
                "observation": [sent("x")],
                "suggestions": [],
            },
        ]
    }
    probs = R.validate_combine(out, {"a", "b"}, per, {"M1", "M2", "M3"}, TEXTS, RESEARCH)
    assert any("must be met" in p for p in probs)
    assert any("can't be missing" in p for p in probs)


def test_validate_summary():
    good = {
        "takeaways": [{"text": f"Point {i}.", "refs": ["F1"], "cites": ["S1"]} for i in range(3)]
    }
    assert R.validate_summary(good, {"F1", "D1"}, RESEARCH) == []
    probs = R.validate_summary(
        {"takeaways": [{"text": "x", "refs": ["F9"], "cites": []}]}, {"F1"}, RESEARCH
    )
    assert any("3 to 6" in p for p in probs)
    assert any("unknown 'F9'" in p for p in probs)
    assert any("must cite a research section" in p for p in probs)


# ---------------- end to end with fakes ----------------


class FakeClient:
    def __init__(self):
        self.calls = []

    def converse(self, **kw):
        tool = kw["toolConfig"]["toolChoice"]["tool"]["name"]
        self.calls.append(tool)
        text = kw["messages"][0]["content"][0]["text"]
        if tool == "record_section":
            out = {
                "findings": [
                    {
                        "question_id": "phonological-awareness",
                        "seen": "yes",
                        "observation": [sent("The lesson models blending.", ["M2"])],
                        "suggestions": [sent("Give quick feedback while blending.", cites=["S1"])],
                    }
                ]
                + [  # every other question asked: not seen here
                    {"question_id": q, "seen": "no", "observation": [sent("Not here.")]}
                    | {"suggestions": []}
                    for q in re.findall(r"^- ([a-z-]+) \(", text, flags=re.M)
                    if q != "phonological-awareness"
                ],
            }
        elif tool == "record_part":  # check-learning shows in the part holding the exit ticket
            mats = re.findall(r'<material id="(M[0-9]+)">\n([^<]*)', text)
            hit = [m for m, t in mats if "Exit ticket" in t]
            yes = "question doc-check-learning" in text and hit
            out = {"answer": "full" if yes else "none", "material": hit if yes else [], "note": ""}
        elif tool == "record_passages":  # the search finds the exit ticket, for check-learning
            ids = re.findall(r'<material id="(M[0-9]+)">\n([^<]*)', text)
            hit = "check whether students learned" in text
            out = {"passages": [m for m, t in ids if hit and "Exit ticket" in t]}
        elif tool == "record_check":  # the cited exit ticket passes the evidence check
            ids = re.findall(r'<material id="(M[0-9]+)">\n([^<]*)', text)
            out = {
                "passages": [
                    {"id": m, "level": "full" if "Exit ticket" in t else "no"} for m, t in ids
                ]
            }
        elif tool == "record_trail":  # quotes the first words of each given part and section
            mats = re.findall(r'<material id="(M[0-9]+)">\n([^<]*)', text)
            secs = re.findall(r'<section id="(S[0-9]+)">\n([^<]*)', text)
            missing = "Verdict: missing" in text
            out = {
                "evidence": [
                    {
                        "material": m,
                        "phrase": " ".join(t.split()[:4]),
                        "shows": "It names the task.",
                    }
                    for m, t in mats[:2]
                ],
                "gaps": [{"looked_for": "a stated objective", "shows": "No outcome is named."}]
                if missing
                else [],
                "research": {"section": secs[0][0], "phrase": " ".join(secs[0][1].split()[:6])}
                if secs
                else {"section": "", "phrase": ""},
                "conclusion": "The research calls for checks; the unit compares as described.",
                "improvements": [
                    {
                        "section": secs[0][0],
                        "phrase": " ".join(secs[0][1].split()[:6]),
                        "apply": "Add a short check at the end of the lesson.",
                        "where": "End of lesson",
                    }
                ]
                if missing and secs
                else [],
            }
        elif tool == "record_grades":
            out = {
                "items": [
                    {"id": i, "grade": "strong"} for i in re.findall(r'<item id="(\w+)">', text)
                ]
            }
        elif tool == "record_verdicts":
            cited = re.findall(r'<material id="(M[0-9]+)">\n[^<]*Exit ticket', text)
            out = {
                "verdicts": [
                    {
                        "question_id": "doc-objective",
                        "verdict": "missing",
                        "observation": [sent("No objective is stated.")],
                        "suggestions": [sent("State the goal at the start.", cites=["S2"])],
                    },
                    {
                        "question_id": "doc-alignment",
                        "verdict": "missing",
                        "observation": [sent("Fit is unclear.")],
                        "suggestions": [sent("Tie each activity to the goal.", cites=["S2"])],
                    },
                    {
                        "question_id": "doc-check-learning",
                        "explicit": True,
                        "verdict": "met",
                        "observation": [sent("An exit ticket checks.", cited[:1] or ["M3"])],
                        "suggestions": [sent("Keep the exit ticket.", cites=["S2"])],
                    },
                ]
            }
            asked = [
                v
                for v in out["verdicts"]
                if f"{v['question_id']} (" in text or f"question {v['question_id']}:" in text
            ]
            out = {"verdicts": asked}  # one call per question
        else:
            out = {
                "takeaways": [
                    {"text": f"Takeaway {i}.", "refs": ["F1", "D1"], "cites": ["S1"]}
                    for i in range(3)
                ]
            }
        return {"output": {"message": {"content": [{"toolUse": {"name": tool, "input": out}}]}}}


def test_run_end_to_end_with_fakes(monkeypatch):
    monkeypatch.setenv("ANSWER_MODEL_ID", "m")
    checklist = R.load_checklist(Path("rubric/checklist.json"))
    material = {"parts": PARTS, "sections": R.make_sections(PARTS)}
    goal = {
        "material_type": "lesson plan",
        "grade_band": "K",
        "focus": "blending",
        "objective": "blend CVC words",
    }

    def embed(text):  # every question looks relevant to every section
        return [1.0, 0.0]

    def search(query, goal):
        sid = (
            "doc-a:v:s1"
            if "assess" in query or "objective" in query or "align" in query
            else "doc-b:v:s1"
        )
        txt = RESEARCH["S2"] if sid == "doc-a:v:s1" else RESEARCH["S1"]
        return [
            {"cite_id": "S1", "text": txt, "label": "Pub, 2016, p. 1", "ref": {"section_id": sid}}
        ]

    client = FakeClient()
    res = review_run.run(
        material, goal, set(), checklist, {"embed": embed, "search": search, "client": client}
    )
    # short material = one section, asked one question per call (all 7 fit grade K)
    assert client.calls.count("record_section") == 7
    assert res["sections"][0]["findings"][0]["id"] == "F1"
    assert [d["question_id"] for d in res["document"]] == [
        "doc-objective",
        "doc-alignment",
        "doc-check-learning",
    ]
    # short material: one whole-document call per question, no chunk answers, no fixed rules
    assert "record_part" not in client.calls and res["stats"]["doc_chunks"] == 1
    assert client.calls.count("record_verdicts") == 3
    assert [d["status"] for d in res["document"]] == [None, None, None]
    assert all(d["answers"] == [] for d in res["document"])
    assert [d["verdict"] for d in res["document"]] == ["missing", "missing", "met"]
    assert res["document_ok"] and not res["summary"]["takeaways"]  # summary off (D88)
    rows = {d["question_id"]: d for d in res["document"]}
    obj, learn = rows["doc-objective"]["trail"], rows["doc-check-learning"]["trail"]
    assert obj["gaps"] and obj["improvements"][0]["match"] == "strong"
    assert learn["research"]["span"] and not learn["improvements"]  # met: no improvements
    m = learn["evidence"][0]
    start, end = m["span"]
    assert PARTS[int(m["material"][1:]) - 1]["text"][start:end]  # highlight sliced from the part
    assert set(res["skipped_document_questions"]) == {
        "doc-sequence",
        "doc-support",
        "doc-practice-time",
        "doc-ms-standards",
    }
    assert set(res["research"]) == {"S1", "S2"}


def test_large_material_is_read_in_chunks_then_merged(monkeypatch):
    monkeypatch.setenv("ANSWER_MODEL_ID", "m")
    monkeypatch.setattr(R, "DOC_CHUNK_WORDS", 1000)
    text = "\n\n".join(
        ["Day 1: Blending", para(700), "Day 2: Practice", para(700), "Day 3: Check", para(200)]
        + ["Exit ticket: each student reads five words aloud."]
    )
    parts = R.material_parts(text)
    material = {"parts": parts, "sections": R.make_sections(parts)}
    goal = {"material_type": "unit", "grade_band": "K", "focus": "x", "objective": "blend words"}

    def search(query, goal):
        return [
            {"cite_id": "S1", "text": RESEARCH[s], "label": "P", "ref": {"section_id": s}}
            for s in ("S1", "S2")
        ]

    client = FakeClient()
    deps = {"embed": lambda t: [1.0], "search": search, "client": client}
    res = review_run.run(material, goal, set(), R.load_checklist("rubric/checklist.json"), deps)
    assert res["stats"]["doc_chunks"] == 2
    assert client.calls.count("record_part") == 2 * 3  # 2 chunks x 3 core questions
    rows = {d["question_id"]: d for d in res["document"]}
    learn = rows["doc-check-learning"]
    assert [a["answer"] for a in learn["answers"]] == ["none", "full"]
    assert learn["verdict"] == "met" and learn["observation"][0]["material"]
    assert rows["doc-alignment"]["status"] == "missing"  # 'all' rule over chunk answers
    assert res["document_ok"]


def test_doc_chunks_pack_whole_sections():
    parts = [{"id": f"M{i}", "text": para(300)} for i in range(1, 7)]
    secs = [
        {"id": f"sec{k}", "title": f"T{k}", "parts": [f"M{i}"], "owned": [f"M{i}"]}
        for k, i in enumerate(range(1, 7), 1)
    ]
    assert R.doc_chunks(parts, secs, 25000) == [
        {"id": "doc1", "title": "Whole material", "parts": [p["id"] for p in parts]}
    ]
    chunks = R.doc_chunks(parts, secs, 700)
    assert [c["parts"] for c in chunks] == [["M1", "M2"], ["M3", "M4"], ["M5", "M6"]]
    assert chunks[0]["title"] == "T1 … T2"


def test_verdict_is_capped_at_three_sentences():
    per = {"a": {"status": None, "answers": [], "single": True}}
    out = {
        "verdicts": [
            {
                "question_id": "a",
                "explicit": True,
                "verdict": "partly",
                "observation": [
                    sent(f"Day {i} has students read words.", ["M2"]) for i in range(4)
                ],
                "suggestions": [sent("Add a check.", cites=["S2"])],
            }
        ]
    }
    probs = R.validate_combine(out, {"a"}, per, {"M1", "M2", "M3"}, TEXTS, RESEARCH)
    assert any("has 4 sentences" in p for p in probs)
    out["verdicts"][0]["observation"] = out["verdicts"][0]["observation"][:2]
    assert R.validate_combine(out, {"a"}, per, {"M1", "M2", "M3"}, TEXTS, RESEARCH) == []


# ---------------- temp store and PowerPoint ----------------


def test_temp_store_purges_after_24_hours(tmp_path):
    store = TempStore(tmp_path, hours=24)
    old, new = store.new(), store.new()
    store.put_json(old, "material.json", {"x": 1})
    past = time.time() - 25 * 3600
    os.utime(tmp_path / old, (past, past))
    assert store.purge() == 1
    assert not (tmp_path / old).exists() and (tmp_path / new).exists()
    with pytest.raises(KeyError):
        store.path(old, "material.json")


def test_pptx_extraction():
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "Blending CVC words"
    s.placeholders[1].text = "Model: c-a-t, cat"
    s.notes_slide.notes_text_frame.text = "Keep it short."
    s2 = prs.slides.add_slide(prs.slide_layouts[5])
    s2.shapes.title.text = "Check"
    t = s2.shapes.add_table(1, 2, Inches(1), Inches(2), Inches(4), Inches(1)).table
    t.cell(0, 0).text, t.cell(0, 1).text = "map", "yes"
    buf = io.BytesIO()
    prs.save(buf)
    ex = extract(buf.getvalue())
    assert ex.kind == "pptx"
    assert ex.text.split("\n\n") == [
        "Blending CVC words",
        "Model: c-a-t, cat",
        "Speaker notes: Keep it short.",
        "Check",
        "map | yes",
    ]
    assert ex.headings == {"Blending CVC words", "Check"}


def test_heading_on_the_line_above_a_paragraph_starts_a_section():
    body = "\n\n".join(f"Day {d}: Practice\n" + para(420) for d in range(1, 5))
    parts = R.material_parts(body)
    assert all(body[p["start"] : p["end"]] == p["text"] for p in parts)
    assert parts[0]["text"] == "Day 1: Practice"
    assert [s["title"] for s in R.make_sections(parts)] == [
        f"Day {d}: Practice" for d in range(1, 5)
    ]


def test_with_retry_survives_a_bedrock_error():
    calls = []

    def call(fb):
        calls.append(fb)
        if len(calls) == 1:
            raise RuntimeError("ModelErrorException: invalid sequence as part of ToolUse")
        return {"ok": 1}

    res = R.with_retry(call, lambda o: [])
    assert res["ok"] and len(calls) == 2 and "invalid sequence" in calls[1]


def test_sanitize_section_fixes_what_is_safe():
    reused = "Teach students to blend individual sounds into words using short daily routines."
    raw = {
        "findings": [
            {
                "question_id": "made-up",
                "seen": "yes",
                "observation": [sent("x", ["M1"])],
                "suggestions": [],
            },
            {
                "question_id": "phonological-awareness",
                "seen": "no",
                "observation": [],
                "suggestions": [],
            },
            {
                "question_id": "phonological-awareness",
                "seen": "yes",
                "observation": [sent("The lesson models blending.", ["M2"])],
                "suggestions": [sent(reused)],
            },
        ],
    }
    clean, notes = R.sanitize_section(raw, {"phonological-awareness"}, SECTION, RESEARCH, TEXTS)
    assert len(clean["findings"]) == 1
    assert clean["findings"][0]["suggestions"][0]["cites"] == ["S1"]  # reuse gets its citation
    assert validate(clean) == []
    assert len(notes) == 2


def test_combine_must_cover_every_question():
    per = {
        "a": {"status": "met", "answers": [{"answer": "full", "material": ["M1"]}]},
        "b": {"status": "missing", "answers": []},
    }
    out = {
        "verdicts": [
            {
                "question_id": "a",
                "explicit": True,
                "verdict": "met",
                "observation": [sent("ok", ["M1"])],
                "suggestions": [sent("keep it", cites=["S1"])],
            }
        ]
    }
    probs = R.validate_combine(out, {"a", "b"}, per, {"M1"}, TEXTS, RESEARCH)
    assert probs == ["write an entry for every question, including fixed verdicts; missing: b"]


def test_unanswered_sections_are_left_out_of_the_combine():
    assert R.combine_status("all", ["full", "unanswered", "full"]) == "met"
    assert R.combine_status("any", ["unanswered"]) == "missing"


def test_borderline_not_seen_finding_is_dropped():
    f = {
        "question_id": "phonological-awareness",
        "seen": "no",
        "observation": [sent("No sound work here.")],
        "suggestions": [],
    }
    for label, kept in (("maybe", 0), ("in", 1)):
        clean, _ = R.sanitize_section(
            {"findings": [dict(f)]},
            {"phonological-awareness"},
            SECTION,
            RESEARCH,
            TEXTS,
            {"phonological-awareness": label},
        )
        assert len(clean["findings"]) == kept


def test_summary_rejects_ids_written_in_text():
    bad = {
        "takeaways": [
            {"text": f"Point {i}, see F1 and S1.", "refs": ["F1"], "cites": ["S1"]}
            for i in range(3)
        ]
    }
    probs = R.validate_summary(bad, {"F1"}, RESEARCH)
    assert len([p for p in probs if "writes IDs in its text" in p]) == 3


def test_strip_ids_removes_citation_talk_only():
    t = (
        "Students practice blending daily. This is supported by findings F1, F2, and F5 "
        "and research sections S3 and S12."
    )
    assert R.strip_ids(t) == "Students practice blending daily."
    assert R.strip_ids("Add a quick check at the end (F4, D2).") == "Add a quick check at the end."
    assert R.strip_ids("Keep the routine short.") == "Keep the routine short."


def test_sanitize_summary_then_validates():
    out = {
        "takeaways": [
            {"text": f"Point {i}. As seen in F1 and S1.", "refs": ["F1"], "cites": ["S1"]}
            for i in range(3)
        ]
    }
    assert R.validate_summary(R.sanitize_summary(out, RESEARCH), {"F1"}, RESEARCH) == []


def test_sanitize_combine_fixes_a_misspelled_single_id():
    out = {
        "verdicts": [
            {
                "question_id": "doc-ms standards",
                "explicit": True,
                "verdict": "met",
                "observation": [],
                "suggestions": [],
            }
        ]
    }
    assert (
        R.sanitize_combine(out, {}, {}, "doc-ms-standards")["verdicts"][0]["question_id"]
        == "doc-ms-standards"
    )


# ---------------- review fixes (2026-10-08) ----------------


def test_word_heading_levels_group_by_top_heading_and_name_the_parent():
    text = "\n\n".join(
        [
            "Unit title",
            para(30),
            "Overview",
            para(100),
            "Day 1",
            "Teach",
            para(500),
            "Practice",
            para(500),
            "Day 2",
            "Warm-up",
            para(40),
            "Badge",
            para(30),
            "Day 3",
            para(700),
        ]
    )
    levels = {"Overview": 1, "Day 1": 1, "Day 2": 1, "Day 3": 1}
    levels |= {"Teach": 2, "Practice": 2, "Warm-up": 2, "Badge": 2}
    parts = R.material_parts(text)
    secs = R.make_sections(parts, set(levels), levels)
    titles = [s["title"] for s in secs]
    assert titles == [
        "Unit title / Overview",  # a tiny opening joins the first top-level group
        "Day 1",
        "Day 1 / Practice",  # over the cap: split at a subheading, parent in the title
        "Day 2",  # small groups never merge across a top-level heading
        "Day 3",
    ]
    by_title = {s["title"]: s["parts"] for s in secs}
    assert all(parts[int(m[1:]) - 1]["text"] != "Day 3" for m in by_title["Day 2"])
    assert [m for s in secs for m in s["parts"]] == [p["id"] for p in parts]


def test_docx_records_heading_levels():
    import docx

    d = docx.Document()
    d.add_heading("Day 1", level=1)
    d.add_heading("Teach", level=2)
    d.add_paragraph("Model the sound.")
    buf = io.BytesIO()
    d.save(buf)
    ex = extract(buf.getvalue())
    assert ex.heading_levels == {"Day 1": 1, "Teach": 2}


def test_goal_is_labeled_as_not_part_of_the_material():
    goal = {"objective": "read sh words", "grade_band": "1", "focus": "x", "material_type": "y"}
    prompt = R.section_prompt(goal, [SECTION], SECTION, {p["id"]: p for p in PARTS}, [], {}, [])
    assert "NOT part of the material" in prompt
    assert "NOT part of" in R.combine_prompt(goal, [], {}, [])
    q = {"id": "doc-objective", "question": "Is there an objective?"}
    assert "NOT part of" in R.doc_prompt(goal, [SECTION], PARTS, q, [])
    assert "NOT part of" in R.part_prompt(goal, [SECTION], SECTION, PARTS, q)


def test_observations_must_name_evidence_not_restate_the_question():
    q = {"phonological-awareness": "Does this part help students hear and work with the sounds?"}
    f = section_out()["findings"][0]
    restated = sent("This part helps students hear and work with the sounds in words.", ["M2"])
    probs = R.validate_section(
        section_out(findings=[{**f, "observation": [f["observation"][0], restated]}]),
        SECTION,
        {"phonological-awareness"},
        TEXTS,
        RESEARCH,
        q,
    )
    assert any("restates the question" in p for p in probs)
    with_evidence = restated["text"].rstrip(".") + " by having partners sort picture cards aloud."
    assert not R.restates(with_evidence, q["phonological-awareness"])
    # after a failed retry, only the restating sentence goes; a finding left empty goes too
    both = section_out(findings=[{**f, "observation": [f["observation"][0], restated]}])
    res = R.with_retry(
        lambda fb: both,
        lambda o: R.validate_section(o, SECTION, {"phonological-awareness"}, TEXTS, RESEARCH, q),
        lambda o: R.drop_restated(o, q),
    )
    assert res["ok"] and res["result"]["findings"][0]["observation"] == [f["observation"][0]]
    only = section_out(findings=[{**f, "observation": [restated]}])
    kept, notes = R.drop_restated(only, q)
    assert kept["findings"] == [] and any("only restated" in n for n in notes)
    unpointed = [f["observation"][0], sent("It also repeats words.")]
    probs = validate(section_out(findings=[{**f, "observation": unpointed}]))
    assert any("observation 1 must point to the material" in p for p in probs)


def test_met_or_partly_verdict_must_rest_on_reported_material():
    per = {
        "a": {"status": None, "answers": [{"answer": "full", "material": ["M3"]}]},
        "b": {"status": None, "answers": [{"answer": "none", "material": []}]},
    }

    def verdict(qid, v, mats):
        return {
            "question_id": qid,
            "explicit": v != "missing",
            "verdict": v,
            "observation": [sent("The exit ticket has students read five words.", mats)],
            "suggestions": [sent("Keep it short.", cites=["S2"])],
        }

    def check(*vs):
        out = {"verdicts": list(vs)}
        return R.validate_combine(out, {"a", "b"}, per, {"M1", "M3"}, TEXTS, RESEARCH)

    assert check(verdict("a", "met", ["M3"]), verdict("b", "missing", [])) == []
    probs = check(verdict("a", "met", ["M1"]), verdict("b", "missing", []))
    assert any("must cite the material it rests on, one of: M3" in p for p in probs)
    probs = check(verdict("a", "met", ["M3"]), verdict("b", "partly", ["M1"]))
    assert any("must be missing: no section showed this" in p for p in probs)


def test_shipped_checklist_judges_objective_check_learning_and_support():
    rules = {
        c["criterion_id"]: c.get("combine")
        for c in R.load_checklist(Path("rubric/checklist.json"))["criteria"]
    }
    assert rules["doc-objective"] == rules["doc-check-learning"] == rules["doc-support"] == "judge"


def test_tool_sequence_error_falls_back_to_json_text(monkeypatch):
    monkeypatch.setenv("ANSWER_MODEL_ID", "m")

    class Err(Exception):
        response = {"Error": {"Code": "ModelErrorException"}}

    class Client:
        def __init__(self):
            self.kw = []

        def converse(self, **kw):
            self.kw.append(kw)
            if "toolConfig" in kw:
                raise Err("Model produced invalid sequence as part of ToolUse.")
            body = '{"answer": "none", "material": [], "note": ""}'
            reply = f"<thinking>hm</thinking>```json\n{body}\n```"
            return {"output": {"message": {"content": [{"text": reply}]}}}

    client = Client()
    res = R.with_retry(
        lambda fb: R.converse(client, "sys", "q", "record_part", R.PART_SCHEMA, 500),
        lambda o: R.validate_part(o, {"parts": ["M1"]}),
    )
    assert res["ok"] and res["result"] == {"answer": "none", "material": [], "note": ""}
    assert res["attempts"] == [{"problems": [], "fallback": "json-text"}]
    assert "JSON schema" in client.kw[1]["messages"][0]["content"][0]["text"]
    assert client.kw[0]["additionalModelRequestFields"] == {"inferenceConfig": {"topK": 1}}

    class Other(Exception):
        response = {"Error": {"Code": "ThrottlingException"}}

    class Throttled:
        def converse(self, **kw):
            raise Other("slow down")

    with pytest.raises(Other):  # only the tool-sequence error falls back
        R.converse(Throttled(), "sys", "q", "record_part", R.PART_SCHEMA, 500)


def test_partly_needs_the_material_to_explicitly_do_part_of_it():
    per = {"a": {"status": None, "answers": [], "single": True, "explicit": True}}
    v = {
        "question_id": "a",
        "explicit": False,
        "verdict": "partly",
        "observation": [sent("Partners build words, which suggests some checking.", ["M2"])],
        "suggestions": [sent("Add an exit ticket.", cites=["S2"])],
    }
    probs = R.validate_combine({"verdicts": [v]}, {"a"}, per, {"M2"}, TEXTS, RESEARCH)
    assert any("the verdict must be missing" in p for p in probs)
    judged = {"a": {**per["a"], "explicit": False}}  # judgment questions skip this rule
    assert R.validate_combine({"verdicts": [v]}, {"a"}, judged, {"M2"}, TEXTS, RESEARCH) == []
    v |= {"verdict": "missing", "observation": [sent("No step checks what students learned.")]}
    assert R.validate_combine({"verdicts": [v]}, {"a"}, per, {"M2"}, TEXTS, RESEARCH) == []


def test_presence_verdict_follows_confirmed_passages(monkeypatch):
    monkeypatch.setenv("ANSWER_MODEL_ID", "m")
    parts_by_id = {p["id"]: p for p in PARTS}

    class Client:  # search finds M1 and M2; only M1 (the stated objective) passes the check
        def converse(self, **kw):
            tool = kw["toolConfig"]["toolChoice"]["tool"]["name"]
            out = {
                "record_passages": {"passages": ["M2", "M1", "M9"]},
                "record_check": {
                    "passages": [{"id": "M1", "level": "full"}, {"id": "M2", "level": "no"}]
                },
            }[tool]
            return {"output": {"message": {"content": [{"toolUse": {"name": tool, "input": out}}]}}}

    found = R.find_evidence(Client(), "Is there an objective?", PARTS)
    assert found == ["M2", "M1"]  # unknown ids dropped
    graded = R.verify_parts(Client(), "Is there an objective?", parts_by_id, found)
    assert graded == {"M1": "full"}
    conf = list(graded)
    q = {"id": "a", "question": "Is there an objective?", "explicit": True, "confirmed": conf}
    assert "confirmed that these passages explicitly do this: M1" in R.question_line(q)

    per = {"a": {"status": None, "answers": [], "single": True, "confirmed": conf}}

    def verdict(v, mats):
        return {
            "question_id": "a",
            "explicit": v != "missing",
            "verdict": v,
            "observation": [sent("The opening line names the skill students will learn.", mats)],
            "suggestions": [sent("Add a way to check it.", cites=["S2"])],
        }

    def check(v):
        out = R.set_confirmed({"verdicts": [v]}, conf)
        return R.validate_combine(out, {"a"}, per, {"M1", "M2", "M3"}, TEXTS, RESEARCH)

    assert check(verdict("partly", ["M1"])) == []
    assert any(
        "must cite at least one of the passages confirmed" in p
        for p in check(verdict("partly", ["M2"]))
    )
    assert any("can't be missing: passages M1" in p for p in check(verdict("missing", [])))
    none = {"a": {**per["a"], "confirmed": []}}
    out = R.set_confirmed({"verdicts": [verdict("met", ["M2"])]}, [])
    probs = R.validate_combine(out, {"a"}, none, {"M1", "M2"}, TEXTS, RESEARCH)
    assert any("found no passage" in p for p in probs)

    # last resort: a missing verdict despite confirmed passages becomes partly, citing them
    out = R.set_confirmed({"verdicts": [verdict("missing", [])]}, conf)
    fixed, notes = R.salvage_verdicts(out)
    v = fixed["verdicts"][0]
    assert v["verdict"] == "partly" and v["observation"][0]["material"] == ["M1"]
    assert "set to partly" in notes[0]


def test_missing_verdicts_keep_only_quoted_parts_and_need_a_suggestion():
    v = {
        "question_id": "a",
        "explicit": False,
        "verdict": "missing",
        "observation": [sent("No step checks learning.", ["M1", "M2"])],
        "suggestions": [],
    }
    out = R.sanitize_combine({"verdicts": [v]}, RESEARCH, {p["id"]: p["text"] for p in PARTS}, "a")
    assert out["verdicts"][0]["observation"][0]["material"] == []  # no chips pointing at nothing
    per = {"a": {"status": None, "answers": [], "single": True}}
    probs = R.validate_combine(out, {"a"}, per, {"M1", "M2", "M3"}, TEXTS, RESEARCH)
    assert any("give at least one suggestion" in p for p in probs)
    kept, notes = R.salvage_verdicts(out)  # last resort keeps the verdict without one
    assert R.validate_combine(kept, {"a"}, per, {"M1", "M2", "M3"}, TEXTS, RESEARCH) == []


def test_section_sentences_must_not_open_with_the_question():
    q = {"phonological-awareness": "Does this part help students hear and work with the sounds?"}
    f = section_out()["findings"][0]
    opener = sent(
        "This part helps students hear and work with the sounds by tapping c-a-t and map.", ["M2"]
    )
    out = section_out(findings=[{**f, "observation": [opener]}])
    probs = R.validate_section(out, SECTION, {"phonological-awareness"}, TEXTS, RESEARCH, q)
    assert any("opens with the question's words" in p for p in probs)
    kept, notes = R.drop_restated(out, q)  # last resort keeps the finding
    assert kept["findings"] and "kept 1 sentence" in notes[-1]
    assert R.validate_section(kept, SECTION, {"phonological-awareness"}, TEXTS, RESEARCH, q) == []


def test_presence_questions_hide_the_goal_and_use_the_strict_scale():
    goal = {"objective": "read sh words", "grade_band": "1", "focus": "x", "material_type": "y"}
    pres = {"id": "doc-objective", "question": "Is there an objective?", "explicit": True}
    judg = {"id": "doc-alignment", "question": "Do activities fit the goal?", "explicit": False}
    p1 = R.doc_prompt(goal, [SECTION], PARTS, pres, [])
    p2 = R.doc_prompt(goal, [SECTION], PARTS, judg, [])
    assert "read sh words" not in p1 and "Never partly because" in p1
    assert "read sh words" in p2 and R.JUDGMENT_SCALE in p2
    assert "Never partial because" in R.part_prompt(goal, [SECTION], SECTION, PARTS, pres)
    bad = {
        "criteria": [
            {**c, "evidence": "maybe"}
            for c in R.load_checklist("rubric/checklist.json")["criteria"][-1:]
        ]
    }
    assert any("kind must be one of" in p for p in R.check_checklist(bad))


def test_salvage_trims_over_cited_verdicts():
    sents = [sent("Days 1 to 7 build words.", [f"M{i}" for i in range(1, 8)])]
    v = {"question_id": "a", "verdict": "met", "observation": sents, "suggestions": "none"}
    out, notes = R.salvage_verdicts({"verdicts": [dict(v)]})
    assert out["verdicts"][0]["observation"][0]["material"] == ["M1", "M2", "M3", "M4", "M5"]
    assert out["verdicts"][0]["suggestions"] == [] and "cited 7 parts" in notes[0]
    checked = {**v, "observation": [sent("x", [f"M{i}" for i in range(1, 8)])], R.VERIFIED: ["M6"]}
    out, _ = R.salvage_verdicts({"verdicts": [checked]})
    assert out["verdicts"][0]["observation"][0]["material"] == ["M6"]  # only confirmed evidence


def test_presence_check_may_overrule_lenient_part_answers():
    answers = [{"section": "Part 1", "answer": "partial", "material": ["M2"], "note": ""}]
    v = {
        "question_id": "a",
        "explicit": False,
        "verdict": "missing",
        "observation": [sent("No step checks what students learned.")],
        "suggestions": [sent("Add an exit ticket.", cites=["S2"])],
    }
    for explicit, ok in ((True, True), (False, False)):
        per = {"a": {"status": None, "answers": answers, "explicit": explicit}}
        probs = R.validate_combine({"verdicts": [v]}, {"a"}, per, {"M2"}, TEXTS, RESEARCH)
        assert (probs == []) is ok


def test_part_answers_cite_at_most_five_parts():
    chunk = {"parts": [f"M{i}" for i in range(1, 10)]}
    out = {"answer": "full", "material": [f"M{i}" for i in (9, 1, 2, 3, 4, 5, 6)], "note": ""}
    assert any("lists 7 parts" in p for p in R.validate_part(out, chunk))
    kept, notes = R.trim_part(out)
    assert kept["material"] == ["M1", "M2", "M3", "M4", "M5"] and R.validate_part(kept, chunk) == []


def test_section_repairs_a_non_list_suggestions_field():
    f = {**section_out()["findings"][0], "suggestions": "Add feedback."}
    clean, _ = R.sanitize_section(
        {"findings": [f]}, {"phonological-awareness"}, SECTION, RESEARCH, TEXTS
    )
    assert clean["findings"][0]["suggestions"] == [] and validate(clean) == []


def test_verdict_trim_keeps_quoted_parts():
    texts = {
        f"M{i}": f"Students read word list number {i} aloud to a partner every day."
        for i in range(1, 8)
    }
    quote = sent("On Day 7 students read word list number 7 aloud to a partner every day.", ["M7"])
    others = sent("Lists recur across the week.", [f"M{i}" for i in range(1, 7)])
    v = {"question_id": "a", "verdict": "met", "observation": [quote, others], "suggestions": []}
    out, _ = R.salvage_verdicts({"verdicts": [v]}, texts)
    kept = {m for x in out["verdicts"][0]["observation"] for m in x["material"]}
    assert "M7" in kept and len(kept) == R.MAX_EVIDENCE


def test_summary_salvage_keeps_cited_takeaways():
    good = [{"text": f"Point {i}.", "refs": ["F1"], "cites": ["S1"]} for i in range(4)]
    bad = [{"text": "Uncited.", "refs": ["F1"], "cites": []}]
    out, notes = R.salvage_summary({"takeaways": good[:2] + bad + good[2:]})
    assert len(out["takeaways"]) == 4 and notes == ["kept 4 of 5 takeaways"]
    assert R.validate_summary(out, {"F1"}, RESEARCH) == []


def test_verdict_needs_an_explanation_and_salvage_supplies_the_standard_one():
    v = {"question_id": "a", "explicit": False, "verdict": "missing", "observation": []}
    v["suggestions"] = [sent("Add a plan for extra help.", cites=["S2"])]
    per = {"a": {"status": None, "answers": [], "single": True}}
    probs = R.validate_combine({"verdicts": [v]}, {"a"}, per, {"M1"}, TEXTS, RESEARCH)
    assert any("needs at least one observation" in p for p in probs)
    out, _ = R.salvage_verdicts({"verdicts": [v]})
    assert out["verdicts"][0]["observation"][0]["text"] == R.NOT_FOUND
    assert R.validate_combine(out, {"a"}, per, {"M1"}, TEXTS, RESEARCH) == []


def test_section_must_answer_every_question_until_the_last_resort():
    raw = section_out()  # answers phonological-awareness only
    qa = {"phonological-awareness", "comprehension"}
    clean, _ = R.sanitize_section(raw, qa, SECTION, RESEARCH, TEXTS)
    probs = R.validate_section(clean, SECTION, qa, TEXTS, RESEARCH)
    assert probs == [
        "answer every section question (seen yes, partly, or no); missing: comprehension"
    ]
    kept, notes = R.drop_restated(clean, {q: "?" for q in qa})
    assert (
        "not answered: comprehension" in notes
        and R.validate_section(kept, SECTION, qa, TEXTS, RESEARCH) == []
    )


def test_observations_never_carry_research():
    f = section_out()["findings"][0]
    borrowed = sent(
        "Teach students to blend individual sounds into words using short daily routines.", ["M2"]
    )
    # cleanup no longer adds the research citation to an observation, so the copy check catches it
    clean, _ = R.sanitize_section(
        section_out(findings=[{**f, "observation": [borrowed]}]),
        {"phonological-awareness"},
        SECTION,
        RESEARCH,
        TEXTS,
    )
    assert clean["findings"][0]["observation"][0]["cites"] == []
    assert validate(clean)  # rejected: reuses S1 wording
    cited = sent("The lesson blends sounds daily.", ["M2"], ["S1"])
    probs = validate(section_out(findings=[{**f, "observation": [cited]}]))
    assert any("observation 0 cites research (S1)" in p for p in probs)
    # last resort drops just that sentence
    both = section_out(findings=[{**f, "observation": [f["observation"][0], cited]}])
    kept, notes = R.drop_restated(both, {"phonological-awareness": "?"}, RESEARCH)
    assert kept["findings"][0]["observation"] == [f["observation"][0]]
    assert "research wording" in notes[0]


def test_met_needs_a_passage_that_does_all_of_it():
    per = {"a": {"status": None, "answers": [], "single": True, "confirmed": ["M1"], "full": []}}
    v = {
        "question_id": "a",
        "explicit": True,
        "verdict": "met",
        "observation": [sent("The opening line names the skill students will learn.", ["M1"])],
        "suggestions": [],
    }
    out = R.set_confirmed({"verdicts": [v]}, ["M1"], [])
    probs = R.validate_combine(out, {"a"}, per, {"M1"}, TEXTS, RESEARCH)
    assert any("can't be met" in p for p in probs)
    fixed, notes = R.salvage_verdicts(out)
    assert fixed["verdicts"][0]["verdict"] == "partly" and "set to partly" in notes[0]


def test_trail_phrases_must_be_word_for_word_and_research_wording_stays_in_quotes():
    research = {"S1": RESEARCH["S1"]}
    good = {
        "evidence": [
            {"material": "M3", "phrase": "each student reads five words", "shows": "A check."}
        ],
        "gaps": [],
        "research": {"section": "S1", "phrase": "blend individual sounds into words"},
        "conclusion": "The research calls for blending practice; the lesson has it.",
        "improvements": [],
    }
    out = R.sanitize_trail(dict(good), "met", {"M3"})
    assert R.validate_trail(out, TEXTS, research) == []
    bad = R.sanitize_trail(
        {**good, "research": {"section": "S1", "phrase": "blend sounds quickly every morning"}},
        "met",
        {"M3"},
    )
    assert any("not word for word" in p for p in R.validate_trail(bad, TEXTS, research))
    fixed, notes = R.salvage_trail(bad, TEXTS, research)
    assert fixed["research"]["section"] == "" and "research left out" in notes[0]
    copied = {
        **good,
        "conclusion": (
            "Teach students to blend individual sounds into words using short daily routines."
        ),
    }
    probs = R.validate_trail(R.sanitize_trail(copied, "met", {"M3"}), TEXTS, research)
    assert any("reuses wording from S1" in p for p in probs)
    # met verdicts get no improvements; parts that were not given are dropped
    extra = {**good, "improvements": [{"section": "S1", "phrase": "x", "apply": "y", "where": "z"}]}
    extra["evidence"] = good["evidence"] + [{"material": "M9", "phrase": "x", "shows": "y"}]
    clean = R.sanitize_trail(extra, "met", {"M3"})
    assert clean["improvements"] == [] and [e["material"] for e in clean["evidence"]] == ["M3"]
    view = R.trail_view(clean, TEXTS | research)
    s0, e0 = view["evidence"][0]["span"]
    assert TEXTS["M3"][s0:e0] == "each student reads five words"


def test_reference_lists_are_not_research():
    from read.retrieve import is_reference_list

    bib = (
        "Case, L. P., Speece, D. L. (2010). Validation of a supplemental reading intervention. "
        "Journal of Learning Disabilities, 43(5), 402-417. Chambers, B. (2011). Small-group "
        "tutoring. Elementary School Journal. Amendum, S. J. (2011). Classroom-based reading."
    )
    guide = (
        "Teachers can monitor student progress and adjust the assigned text for students of "
        "above- or below-average reading ability. The National Reading Panel (2000) agreed."
    )
    assert is_reference_list(bib) and not is_reference_list(guide)


def test_find_phrase_tolerates_line_breaks_and_quotes():
    t = "Panel’s Advice. Teachers can monitor stu- dent progress and adjust the text."
    assert R.find_phrase(t, "monitor student progress and adjust") == (29, 66)
    assert R.find_phrase(t, "Panel's Advice") == (0, 14)
    assert R.find_phrase(t, "words that are not there") is None


def test_off_topic_research_takes_its_conclusion_with_it(monkeypatch):
    monkeypatch.setenv("ANSWER_MODEL_ID", "m")

    class Client:
        def converse(self, **kw):
            out = {"items": [{"id": "R", "grade": "no"}, {"id": "I0", "grade": "limited"}]}
            tool = kw["toolConfig"]["toolChoice"]["tool"]["name"]
            return {"output": {"message": {"content": [{"toolUse": {"name": tool, "input": out}}]}}}

    trail = {
        "evidence": [],
        "gaps": [],
        "research": {"section": "S1", "phrase": "blend individual sounds"},
        "conclusion": "The research calls for blending.",
        "improvements": [
            {"section": "S2", "phrase": "x", "apply": "Add a check.", "where": "Day 5"}
        ],
    }
    out = R.research_checks(trail, RESEARCH, Client(), "Does it check learning?")
    assert out["research"]["section"] == "" and out["conclusion"] == ""
    assert out["improvements"][0]["match"] == "limited"
